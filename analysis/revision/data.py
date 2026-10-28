"""Reproducible, audited data preparation for ICDEES paper 1571345184.

Monthly reference calculations are approximations at completed month ages, not
exact-day clinical gold standards. Jeneponto lacks measurement position and child
identifiers: no length/height conversion or missing-value estimation is invented.
"""
from __future__ import annotations

from concurrent.futures import ThreadPoolExecutor
from datetime import date, datetime, timezone
import hashlib
import json
import math
from pathlib import Path
import re

import numpy as np
import pandas as pd
import requests
from openpyxl import load_workbook


FEATURES = ['Gender_Encoded', 'Age (months)', 'Height (cm)', 'Weight (kg)']
JENEPONTO_URL = 'https://data.mendeley.com/public-files/datasets/wzwpc9j5bx/files/43dc7b64-13fe-4ab7-ac00-1eb26c98f612/file_downloaded'
JENEPONTO_SHA256 = '4bd5e56aa0839f3ecdb39cc60dfc194a5bb12e0f300b49c8fc7ec0e43d010faf'
WHO_PAGE = 'https://www.who.int/tools/child-growth-standards/standards/length-height-for-age'
WHO_WFA_PAGE = 'https://www.who.int/tools/child-growth-standards/standards/weight-for-age'
WHO_ALGORITHM = 'https://github.com/WorldHealthOrganization/anthro/blob/master/R/z-score-helper.R'
_WHO_BASE = 'https://cdn.who.int/media/docs/default-source/child-growth/child-growth-standards/indicators/'
WHO_FILES = {
    'lfa_boys': _WHO_BASE + 'length-height-for-age/lhfa_boys_0-to-2-years_zscores.xlsx?sfvrsn=30e044c_9',
    'hfa_boys': _WHO_BASE + 'length-height-for-age/lhfa_boys_2-to-5-years_zscores.xlsx?sfvrsn=17e5ad91_9',
    'lfa_girls': _WHO_BASE + 'length-height-for-age/lhfa_girls_0-to-2-years_zscores.xlsx?sfvrsn=e9e66a95_11',
    'hfa_girls': _WHO_BASE + 'length-height-for-age/lhfa_girls_2-to-5-years_zscores.xlsx?sfvrsn=2ec187b9_11',
    'wfa_boys': _WHO_BASE + 'weight-for-age/wfa_boys_0-to-5-years_zscores.xlsx?sfvrsn=97a05331_9',
    'wfa_girls': _WHO_BASE + 'weight-for-age/wfa_girls_0-to-5-years_zscores.xlsx?sfvrsn=4c03b8db_7',
}


def _root(root=None):
    return Path(root).resolve() if root is not None else Path(__file__).resolve().parents[2]


def _write_json(path, data):
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(data, indent=2, ensure_ascii=False, allow_nan=False) + '\n')


def _download(url, path, expected_sha=None):
    if not path.exists():
        response = requests.get(url, timeout=(15, 60))
        response.raise_for_status()
        if not response.content.startswith(b'PK'):
            raise ValueError(f'Expected an XLSX ZIP container from {url}')
        digest = hashlib.sha256(response.content).hexdigest()
        if expected_sha and digest != expected_sha:
            raise ValueError(f'Source checksum changed for {url}: {digest}')
        path.parent.mkdir(parents=True, exist_ok=True)
        path.write_bytes(response.content)
    digest = hashlib.sha256(path.read_bytes()).hexdigest()
    if expected_sha and digest != expected_sha:
        raise ValueError(f'Cached source checksum does not match: {path}')
    return {'url': url, 'sha256': digest, 'bytes': path.stat().st_size,
            'retrieved_at_utc': datetime.fromtimestamp(path.stat().st_mtime, timezone.utc).isoformat(),
            'verified_at_utc': datetime.now(timezone.utc).isoformat()}


def _legacy_references(root):
    directory = root / 'dataset' / 'Stunting Wasting Dataset (Synthetic)' / 'z-scores who'
    frames = []
    for sex, gender in [('boys', 1), ('girls', 0)]:
        for indicator, prefix in [('HAZ', 'lhfa'), ('WAZ', 'wfa')]:
            table = pd.read_csv(directory / f'{prefix}_{sex}_0-to-5-years_zscores.csv', sep=';', decimal=',')
            table = table.rename(columns={table.columns[0]: 'Age (months)'})
            table = table[['Age (months)', 'L', 'M', 'S']].copy()
            table['Indicator'] = indicator
            table['Gender_Encoded'] = gender
            frames.append(table)
    return pd.concat(frames, ignore_index=True)[['Indicator', 'Gender_Encoded', 'Age (months)', 'L', 'M', 'S']]


def prepare_references(root=None):
    """Cache primary WHO XLSX files and audit every original LMS coefficient.

    The canonical table uses length below 24 months and height from 24 months.
    The overlapping 24-month rows are never averaged.
    """
    root = _root(root)
    directory = root / 'dataset' / 'revision' / 'references'
    target = directory / 'who_monthly_corrected.csv'
    if target.exists():
        manifest_path = directory / 'manifest.json'
        if manifest_path.exists():
            manifest = json.loads(manifest_path.read_text())
            expected = manifest.get('corrected_sha256')
            if expected and hashlib.sha256(target.read_bytes()).hexdigest() != expected:
                raise ValueError('Corrected reference cache checksum differs from its manifest')
        return pd.read_csv(target)
    raw = root / 'dataset' / 'revision' / 'raw' / 'who'
    with ThreadPoolExecutor(max_workers=3) as pool:
        entries = list(pool.map(lambda item: (item[0], _download(item[1], raw / f'{item[0]}.xlsx')), WHO_FILES.items()))
    frames = []
    for name, _ in entries:
        table = pd.read_excel(raw / f'{name}.xlsx')
        table.columns = [str(c).strip() for c in table.columns]
        table = table.rename(columns={'Month': 'Age (months)'})
        table = table[['Age (months)', 'L', 'M', 'S']].copy()
        table['Indicator'] = 'WAZ' if name.startswith('wfa') else 'HAZ'
        table['Gender_Encoded'] = 1 if name.endswith('boys') else 0
        if name.startswith('lfa'):
            table = table[table['Age (months)'] < 24]
        elif name.startswith('hfa'):
            table = table[table['Age (months)'] >= 24]
        frames.append(table)
    corrected = pd.concat(frames, ignore_index=True)[['Indicator', 'Gender_Encoded', 'Age (months)', 'L', 'M', 'S']]
    corrected = corrected.sort_values(['Indicator', 'Gender_Encoded', 'Age (months)']).reset_index(drop=True)
    if len(corrected) != 244 or corrected.duplicated(['Indicator', 'Gender_Encoded', 'Age (months)']).any():
        raise ValueError('Unexpected primary WHO reference coverage: expected 244 unique rows')
    comparison = _legacy_references(root).merge(corrected, on=['Indicator', 'Gender_Encoded', 'Age (months)'], suffixes=('_legacy', '_corrected'), validate='one_to_one')
    for coefficient in ['L', 'M', 'S']:
        comparison[f'{coefficient}_Changed'] = ~np.isclose(comparison[f'{coefficient}_legacy'], comparison[f'{coefficient}_corrected'], rtol=0, atol=1e-12)
    comparison['Any_Changed'] = comparison[['L_Changed', 'M_Changed', 'S_Changed']].any(axis=1)
    directory.mkdir(parents=True, exist_ok=True)
    comparison.to_csv(directory / 'legacy_vs_primary_who_coefficients.csv', index=False)
    corrected.to_csv(target, index=False)
    audit = {
        'source_pages': [WHO_PAGE, WHO_WFA_PAGE], 'raw_sources': dict(entries),
        'reference_rows': len(corrected), 'coefficients_compared': len(comparison) * 3,
        'corrected_sha256': hashlib.sha256(target.read_bytes()).hexdigest(),
        'changed_rows': comparison[comparison.Any_Changed].to_dict('records'),
        'convention': 'LFA for completed months 0–23; HFA for completed months 24–60; no measurement-position conversion when unknown.',
        'algorithm_source': WHO_ALGORITHM,
        'waz_tail': 'WHO restricted LMS: extrapolate beyond +/-3 SD using distances between SD2 and SD3 measurements.',
        'limitation': 'Monthly ages approximate exact-day WHO assessment; absent measurement position is an assumption, never observed metadata.',
        'licence_note': 'WHO primary reference tables are attributed to WHO; no third-party licence is invented.',
    }
    _write_json(directory / 'manifest.json', audit)
    return corrected


def monthly_zscore(values, L, M, S, restrict_tails=False):
    """LMS, with an optional WHO measurement-distance extension outside 3 SD."""
    values, L, M, S = np.broadcast_arrays(np.asarray(values, dtype=float), np.asarray(L, dtype=float), np.asarray(M, dtype=float), np.asarray(S, dtype=float))
    with np.errstate(divide='ignore', invalid='ignore', over='ignore'):
        logarithmic = np.log(values / M) / S
        score = np.divide((values / M) ** L - 1, L * S, out=logarithmic.copy(), where=L != 0)
        if restrict_tails:
            def measurement(z):
                return np.where(L == 0, M * np.exp(S * z), M * (1 + L * S * z) ** np.divide(1, L, out=np.zeros_like(L), where=L != 0))
            high3, high2, low3, low2 = (measurement(z) for z in [3, 2, -3, -2])
            score = np.where(score > 3, 3 + (values - high3) / (high3 - high2), score)
            score = np.where(score < -3, -3 + (values - low3) / (low2 - low3), score)
    return np.where((values > 0) & (M > 0) & (S > 0), score, np.nan)


def _enrich(frame, reference, corrected=True):
    result = frame.copy()
    for indicator, measure in [('HAZ', 'Height (cm)'), ('WAZ', 'Weight (kg)')]:
        table = reference[reference.Indicator == indicator].set_index(['Gender_Encoded', 'Age (months)'])
        keys = pd.MultiIndex.from_frame(result[['Gender_Encoded', 'Age (months)']])
        coefficients = table[['L', 'M', 'S']].reindex(keys)
        result[indicator] = monthly_zscore(result[measure].to_numpy(), coefficients.L.to_numpy(), coefficients.M.to_numpy(), coefficients.S.to_numpy(), restrict_tails=corrected and indicator == 'WAZ')
    return result


def _profiles(frame):
    groups = {
        'four_features': FEATURES,
        'sex_age_height': FEATURES[:3],
        'sex_age_weight': ['Gender_Encoded', 'Age (months)', 'Weight (kg)'],
    }
    result = {}
    for name, columns in groups.items():
        counts = frame.groupby(columns, dropna=False).size()
        result[name] = {'unique': int(len(counts)), 'repeat_excess': int(len(frame) - len(counts)),
                        'max_group_size': int(counts.max()) if len(counts) else 0}
    return result


def _category(haz):
    return pd.Series(np.select([haz < -3, haz < -2, haz > 3], ['Severely Stunted', 'Stunted', 'Tall'], default='Normal'), index=haz.index)


def load_synthetic(reference='corrected', root=None):
    """Return original synthetic classification labels with the chosen Z-score targets."""
    if reference not in {'legacy', 'corrected'}:
        raise ValueError("reference must be 'legacy' or 'corrected'")
    root = _root(root)
    path = root / 'dataset' / 'Stunting Wasting Dataset (Synthetic)' / 'stunting_wasting_dataset.csv'
    original = pd.read_csv(path)
    frame = original.rename(columns={'Jenis Kelamin': 'Gender', 'Umur (bulan)': 'Age (months)', 'Tinggi Badan (cm)': 'Height (cm)', 'Berat Badan (kg)': 'Weight (kg)'}).copy()
    frame['Gender_Encoded'] = frame.Gender.map({'Laki-laki': 1, 'Perempuan': 0, 'Male': 1, 'Female': 0})
    if frame['Gender_Encoded'].isna().any():
        raise ValueError('Unknown synthetic gender label')
    frame['Gender_Encoded'] = frame['Gender_Encoded'].astype(int)
    frame['Source_Row'] = np.arange(1, len(frame) + 1)
    table = _legacy_references(root) if reference == 'legacy' else prepare_references(root)
    frame = _enrich(frame, table, corrected=reference == 'corrected')
    frame = frame[FEATURES + ['Stunting', 'HAZ', 'WAZ', 'Source_Row']]
    audit = {'source': str(path.relative_to(root)), 'sha256': hashlib.sha256(path.read_bytes()).hexdigest(),
             'rows': len(frame), 'reference': reference, 'original_classification_labels_preserved': True,
             'class_counts': {str(k): int(v) for k, v in frame.Stunting.value_counts().items()},
             'label_vs_recomputed_haz_mismatches': int((_category(frame.HAZ) != frame.Stunting).sum()),
             'profiles': _profiles(frame), 'age_counts': {str(k): int(v) for k, v in frame['Age (months)'].value_counts().sort_index().items()},
             'finite_haz': int(np.isfinite(frame.HAZ).sum()), 'finite_waz': int(np.isfinite(frame.WAZ).sum())}
    if reference == 'corrected':
        old = _enrich(frame, _legacy_references(root), corrected=False)
        audit['haz_changed_from_legacy'] = int((~np.isclose(frame.HAZ, old.HAZ, rtol=0, atol=1e-12)).sum())
        audit['waz_changed_from_legacy'] = int((~np.isclose(frame.WAZ, old.WAZ, rtol=0, atol=1e-12)).sum())
    _write_json(root / 'dataset' / 'revision' / 'audits' / f'synthetic_{reference}.json', audit)
    return frame, audit


def parse_weight(value, number_format='General'):
    """Return (measurement, restored_display, error); never estimate missing data."""
    if isinstance(value, (datetime, date)):
        if number_format != r'd\.m':
            return math.nan, False, 'unsupported_date_format'
        return float(f'{value.day}.{value.month}'), True, None
    if value is None or isinstance(value, (bool, np.bool_)) or (isinstance(value, str) and not re.fullmatch(r'\s*[+-]?(?:\d+(?:\.\d*)?|\.\d+)\s*', value)):
        return math.nan, False, 'missing_or_malformed_weight'
    try:
        numeric = float(value)
    except (TypeError, ValueError):
        return math.nan, False, 'missing_or_malformed_weight'
    return (numeric, False, None) if math.isfinite(numeric) else (math.nan, False, 'nonfinite_weight')


def parse_jeneponto_workbook(path, reference):
    """Retain every source record, annotate deterministic task-specific validity."""
    workbook = load_workbook(path, read_only=True, data_only=True)
    sheet = workbook.active
    iterator = sheet.iter_rows()
    headers = [str(cell.value).strip() for cell in next(iterator)]
    required = ['No.', 'Gender', 'Age (Month)', 'Weight', 'Height', 'Height for Age']
    if any(name not in headers for name in required):
        raise ValueError('Jeneponto workbook schema does not contain the required source fields')
    records = []
    for excel_row, cells in enumerate(iterator, 2):
        if all(cell.value is None for cell in cells):
            continue
        row = {name: cell.value for name, cell in zip(headers, cells)}
        weight_cell = cells[headers.index('Weight')]
        weight, restored, error = parse_weight(weight_cell.value, weight_cell.number_format)
        record = {
            'Source_Row': row['No.'], 'Source_Excel_Row': excel_row,
            'Gender_Encoded': {'M': 1, 'F': 0}.get(row['Gender'], math.nan),
            'Age (months)': row['Age (Month)'], 'Height (cm)': row['Height'], 'Weight (kg)': weight,
            'Stunting_Binary': {'Not Stunted': 0, 'Stunted': 1}.get(row['Height for Age'], math.nan),
            'Source_Status': row['Height for Age'], 'Weight_Raw': str(weight_cell.value),
            'Weight_Number_Format': weight_cell.number_format, 'Weight_Display_Restored': restored,
            'Weight_Parse_Error': error or '',
        }
        for source_name, output_name in [('Z-Score H/A', 'Source_HAZ'), ('Z-Score  W/A', 'Source_WAZ')]:
            value = row.get(source_name)
            record[output_name] = value if isinstance(value, (int, float, str)) and not isinstance(value, bool) else math.nan
        records.append(record)
    workbook.close()
    frame = pd.DataFrame(records)
    for column in ['Age (months)', 'Height (cm)', 'Source_HAZ', 'Source_WAZ']:
        frame[column] = pd.to_numeric(frame[column], errors='coerce')
    frame = _enrich(frame, reference)
    age = frame['Age (months)']
    frame['Flag_Age'] = ~np.isfinite(age) | ~age.between(0, 24) | age.mod(1).ne(0)
    frame['Flag_Gender'] = ~frame.Gender_Encoded.isin([0, 1])
    frame['Flag_Height'] = ~np.isfinite(frame['Height (cm)']) | frame['Height (cm)'].le(0)
    frame['Flag_Weight'] = ~np.isfinite(frame['Weight (kg)']) | frame['Weight (kg)'].le(0)
    frame['Flag_Source_Label'] = ~frame.Stunting_Binary.isin([0, 1])
    frame['Flag_HAZ'] = ~np.isfinite(frame.HAZ) | ~frame.HAZ.between(-6, 6)
    frame['Flag_WAZ'] = ~np.isfinite(frame.WAZ) | ~frame.WAZ.between(-6, 5)
    basic = ~frame.Flag_Age & ~frame.Flag_Gender
    frame['Valid_HAZ'] = basic & ~frame.Flag_Height & ~frame.Flag_HAZ
    frame['Valid_WAZ'] = basic & ~frame.Flag_Weight & ~frame.Flag_WAZ
    frame['Valid_Classification'] = frame.Valid_HAZ & frame.Valid_WAZ & ~frame.Flag_Source_Label
    frame['In_Primary_External'] = age.ge(0) & age.lt(24)
    frame['In_Age24_Sensitivity'] = age.eq(24)
    frame['Source_vs_Monthly_HAZ_Mismatch'] = frame.Valid_HAZ & ~frame.Flag_Source_Label & frame.Stunting_Binary.ne((frame.HAZ < -2).astype(int))
    primary = frame[frame.Valid_Classification & frame.In_Primary_External]
    comparison = frame[frame.Valid_HAZ & ~frame.Flag_Source_Label]
    audit = {
        'rows': len(frame), 'rows_age_0_to_24': int((~frame.Flag_Age).sum()),
        'rows_age_24': int(frame.In_Age24_Sensitivity.sum()),
        'date_weights_restored': int(frame.Weight_Display_Restored.sum()),
        'weight_parse_errors': {str(k): int(v) for k, v in frame.loc[frame.Weight_Parse_Error.ne(''), 'Weight_Parse_Error'].value_counts().items()},
        'source_label_mismatches': int(frame.Source_vs_Monthly_HAZ_Mismatch.sum()),
        'source_label_comparison_denominator': len(comparison),
        'primary_source_label_mismatches': int((frame.Source_vs_Monthly_HAZ_Mismatch & frame.In_Primary_External).sum()),
        'primary_source_label_comparison_denominator': int((frame.Valid_HAZ & ~frame.Flag_Source_Label & frame.In_Primary_External).sum()),
        'age24_source_label_mismatches': int((frame.Source_vs_Monthly_HAZ_Mismatch & frame.In_Age24_Sensitivity).sum()),
        'flags': {c: int(frame[c].sum()) for c in frame if c.startswith('Flag_')},
        'valid_all_age_0_to_24': {task: int(frame[f'Valid_{task}'].sum()) for task in ['Classification', 'HAZ', 'WAZ']},
        'valid_primary_age_lt_24': {task: int((frame[f'Valid_{task}'] & frame.In_Primary_External).sum()) for task in ['Classification', 'HAZ', 'WAZ']},
        'valid_sensitivity_age_24': {task: int((frame[f'Valid_{task}'] & frame.In_Age24_Sensitivity).sum()) for task in ['Classification', 'HAZ', 'WAZ']},
        'primary_binary_counts': {str(int(k)): int(v) for k, v in primary.Stunting_Binary.value_counts().items()},
        'primary_profiles': _profiles(primary),
        'all_age_0_to_24_profiles': _profiles(frame[frame.Valid_Classification]),
        'age_groups': {name: int(mask.sum()) for name, mask in {'0–5': age.between(0, 5), '6–11': age.between(6, 11), '12–17': age.between(12, 17), '18–23': age.between(18, 23), '24': age.eq(24)}.items()},
        'reference': 'Primary WHO monthly LMS: length below 24 months, height at 24; restricted WAZ tails.',
        'limitations': ['No unique child identifier: No. is a source row ordinal; repeated profiles are not proven repeat visits.',
                        'Measurement position is unavailable; values are used as recorded, with age-based reference assumptions and no 0.7 cm conversion.',
                        'Monthly recalculated Z-scores are not exact-day clinical gold standards; supplied binary labels are preserved.',
                        'A single regional dataset does not establish national representativeness or prospective prediction.'],
    }
    return frame, audit


def load_jeneponto(root=None):
    """Download/cache the immutable 2024 cohort and return all source rows + QA masks."""
    root = _root(root)
    directory = root / 'dataset' / 'revision'
    raw = directory / 'raw' / 'jeneponto_2024.xlsx'
    manifest = _download(JENEPONTO_URL, raw, JENEPONTO_SHA256)
    manifest.update({'doi': '10.17632/wzwpc9j5bx.4', 'landing_page': 'https://data.mendeley.com/datasets/wzwpc9j5bx/4',
                     'title': 'Dataset Stunting and Nutritional Status of Toddler from Jeneponto Regency, South Sulawesi, Indonesia',
                     'creators': ['Asmaul Husna RS', 'Purnamawati Purnamawati', 'Hendra Jaya', 'Anas Arfandi', 'Iwan Suhardi', 'Slamet Widodo', 'Syahrani Lonang'],
                     'licence': 'CC BY 4.0', 'cohort': 'Jeneponto 2024, source measurements; not a national sample'})
    _write_json(directory / 'raw' / 'jeneponto_manifest.json', manifest)
    frame, audit = parse_jeneponto_workbook(raw, prepare_references(root))
    audit['source'] = manifest
    derived = directory / 'derived'
    derived.mkdir(parents=True, exist_ok=True)
    frame.to_csv(derived / 'jeneponto_2024_harmonized.csv', index=False)
    _write_json(directory / 'audits' / 'jeneponto.json', audit)
    return frame, audit


if __name__ == '__main__':
    prepare_references()
    _, old = load_synthetic('legacy')
    _, corrected = load_synthetic('corrected')
    _, external = load_jeneponto()
    print(json.dumps({'synthetic_legacy': old, 'synthetic_corrected': corrected, 'external': external}, indent=2))
